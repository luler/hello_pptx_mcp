# -*- coding: utf-8 -*-
"""FastAPI 应用：同时提供 MCP（streamable HTTP）与 REST 接口。

    uvicorn app:app --host 0.0.0.0 --port 48000

管理后台：
    GET  /admin              现代化 PPTX 资产在线管理后台 (SPA)
    GET  /manage             管理后台别名
    POST /api/admin/verify   校验 AUTH_KEY 授权
    GET  /api/admin/files    已生成文稿列表（按时间倒序）
    POST /api/admin/delete   批量彻底删除文稿及预览图

REST 端点（便于 curl / 前端直接调用）：
    GET  /                   服务与端点说明
    GET  /healthz            健康检查
    GET  /api/themes         主题列表
    GET  /api/templates      模板列表
    GET  /api/templates/{n}  模板 spec
    POST /api/validate       校验 spec
    POST /api/generate       生成 PPTX（自适应 Host，支持 preview 与密度体检）
    POST /api/artifacts/{id}/page/{n} 单页精修
    GET  /api/artifacts      产物列表
    GET  /api/artifacts/{id} 产物元信息
    GET  /api/download/{id}  下载产物（自适应 filename / id）
    GET  /api/preview/{id}   预览 PNG（?page=1）
    DELETE /api/artifacts/{id}  删除产物
    /mcp                     MCP Streamable HTTP 端点
"""
from __future__ import annotations

import contextlib
import glob
import os
import shutil
import time
from urllib.parse import parse_qs

from fastapi import FastAPI, Header, HTTPException, Query, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse
from pydantic import BaseModel, Field

import mcp_server as M
from pptkit import theme as theme_mod, templates
from pptkit.deck import build_from_spec
from web_admin import get_admin_html

# 读取环境变量中的授权密钥
AUTH_KEY = os.environ.get("AUTH_KEY", "").strip()


def check_auth(authorization: str | None = None, auth_key: str | None = None) -> bool:
    """若服务端配置了 AUTH_KEY，验证 Authorization 请求头或 URL Query 参数 auth_key。"""
    if not AUTH_KEY:
        return True  # 未设置密钥时全量放行（向下兼容）

    # 1. 检查 Authorization: Bearer <AUTH_KEY>
    if authorization:
        parts = authorization.strip().split()
        if len(parts) == 2 and parts[0].lower() == "bearer":
            if parts[1] == AUTH_KEY:
                return True
        elif len(parts) == 1 and parts[0] == AUTH_KEY:
            return True

    # 2. 检查 Query 参数 ?auth_key=<AUTH_KEY>（用于 <img> 预览和 <a> 下载等无法自定义 Header 的场景）
    if auth_key and auth_key == AUTH_KEY:
        return True

    return False


def _delete_presentation_and_preview(item_id: str) -> bool:
    """彻底删除演示文稿实体文件及其对应关联的高清预览图缓存目录。"""
    found = False
    item_id = str(item_id or "").strip()
    if not item_id:
        return False

    rec = M.STORE.get(item_id)
    target_path = None
    target_name = None

    if rec:
        target_path = rec.get("path")
        target_name = rec.get("name")
        M.STORE.delete(rec["id"])
        found = True

    if not target_path or not os.path.exists(target_path):
        cand = os.path.join(M.OUTPUT_DIR, item_id)
        if not os.path.exists(cand) and not cand.endswith(".pptx"):
            cand += ".pptx"
        if os.path.exists(cand):
            target_path = cand
            target_name = os.path.basename(cand)

    if target_path and os.path.exists(target_path):
        try:
            os.remove(target_path)
            found = True
        except OSError:
            pass

    # 清理关联的 preview 预览图目录
    name_for_slug = target_name or item_id
    stem = os.path.splitext(name_for_slug)[0]
    slug = M._slug(stem, "deck")

    dirs_to_clean = {
        os.path.join(M.OUTPUT_DIR, "preview", slug),
        os.path.join(M.OUTPUT_DIR, "preview", stem),
        os.path.join(M.OUTPUT_DIR, "preview", item_id),
    }
    if target_path:
        dirs_to_clean.add(os.path.join(os.path.dirname(target_path), "preview", slug))

    for pv in dirs_to_clean:
        if os.path.isdir(pv):
            shutil.rmtree(pv, ignore_errors=True)
            found = True

    return found


class SpecIn(BaseModel):
    spec: dict = Field(..., description="演示文稿定义，需含 pages")
    filename: str = ""
    preview: bool = True
    store: bool = True


class PageUpdateIn(BaseModel):
    page_spec: dict = Field(..., description="更新后的单页 spec 定义（必须含 layout）")


class VerifyIn(BaseModel):
    key: str = ""


class BatchDeleteIn(BaseModel):
    ids: list[str] = Field(default_factory=list, description="待删除的文稿 ID 或文件名列表")


@contextlib.asynccontextmanager
async def lifespan(app: FastAPI):
    # streamable_http_app() 会惰性建立 session_manager，必须在启动前调用
    async with M.server.session_manager.run():
        yield


app = FastAPI(
    title="pptx-studio",
    description="声明式 PPTX 生成服务：FastAPI + MCP",
    version="0.2.0",
    lifespan=lifespan,
)


try:
    from mcp.server.transport_security import TransportSecurityMiddleware, TransportSecuritySettings
    # 彻底放行所有域名、反向代理与 Origin，彻底根治 421 Invalid Host header
    TransportSecurityMiddleware._validate_host = lambda self, host: True
    TransportSecurityMiddleware._validate_origin = lambda self, origin: True
except Exception:
    pass


class BaseUrlMiddleware:
    """纯 ASGI 中间件：
    1. 提取 HTTP 客户端真实请求的协议与 Host 头，注入上下文供 URL 动态拼接；
    2. 解决 /mcp 重定向问题；
    3. 若配置了 AUTH_KEY，拦截并验证 /mcp 端点的 Bearer 授权请求头。
    """
    def __init__(self, inner_app):
        self.inner_app = inner_app

    async def __call__(self, scope, receive, send):
        if scope.get("type") == "http":
            path = scope.get("path", "")
            method = scope.get("method", "GET").upper()

            # 解决客户端请求 /mcp 时被 Starlette 307 重定向的问题（免去 307 跳转）
            if path == "/mcp":
                path = "/mcp/"
                scope["path"] = "/mcp/"

            headers_list = scope.get("headers", [])
            headers = dict(headers_list)
            x_proto = headers.get(b"x-forwarded-proto", b"").decode("latin-1")
            proto = x_proto or scope.get("scheme", "http")
            x_host = headers.get(b"x-forwarded-host", b"").decode("latin-1")
            host_header = headers.get(b"host", b"").decode("latin-1")
            host = x_host or host_header or ""
            if host:
                M.current_request_base_url.set(f"{proto}://{host}")

            # 若配置了 AUTH_KEY，对 /mcp 及其所有子路由强制要求 Bearer 授权（放行 OPTIONS 跨域预检）
            if AUTH_KEY and (path == "/mcp" or path.startswith("/mcp/")):
                if method != "OPTIONS":
                    auth_header = headers.get(b"authorization", b"").decode("latin-1").strip()
                    authorized = False
                    if auth_header:
                        parts = auth_header.split()
                        if len(parts) == 2 and parts[0].lower() == "bearer" and parts[1] == AUTH_KEY:
                            authorized = True
                        elif len(parts) == 1 and parts[0] == AUTH_KEY:
                            authorized = True

                    if not authorized:
                        query_str = scope.get("query_string", b"").decode("latin-1")
                        qs = parse_qs(query_str)
                        if qs.get("auth_key", [None])[0] == AUTH_KEY or qs.get("token", [None])[0] == AUTH_KEY:
                            authorized = True

                    if not authorized:
                        body = b'{"error": "Unauthorized: missing or invalid Bearer token"}\n'
                        await send({
                            "type": "http.response.start",
                            "status": 401,
                            "headers": [
                                (b"content-type", b"application/json"),
                                (b"content-length", str(len(body)).encode("ascii")),
                                (b"www-authenticate", b'Bearer realm="pptx-studio"'),
                                (b"access-control-allow-origin", b"*"),
                            ],
                        })
                        await send({
                            "type": "http.response.body",
                            "body": body,
                        })
                        return

        await self.inner_app(scope, receive, send)


app.add_middleware(BaseUrlMiddleware)
app.add_middleware(
    CORSMiddleware,
    allow_origin_regex=".*",
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# MCP 端点挂在 /mcp（Streamable HTTP）。
_mcp_app = M.server.streamable_http_app(streamable_http_path="/")

try:
    _mgr = getattr(M.server, "session_manager", None) or getattr(M.server, "_session_manager", None)
    if _mgr:
        if hasattr(_mgr, "security_settings") and _mgr.security_settings:
            _mgr.security_settings.enable_dns_rebinding_protection = False
            _mgr.security_settings.allowed_hosts = ["*"]
            _mgr.security_settings.allowed_origins = ["*"]
        if hasattr(_mgr, "_security") and hasattr(_mgr._security, "settings"):
            _mgr._security.settings.enable_dns_rebinding_protection = False
            _mgr._security.settings.allowed_hosts = ["*"]
            _mgr._security.settings.allowed_origins = ["*"]
except Exception:
    pass

app.mount("/mcp", _mcp_app)


# ==================== 管理后台前端页面 ====================

@app.get("/admin", response_class=HTMLResponse)
@app.get("/manage", response_class=HTMLResponse)
def admin_page():
    """现代化 PPTX 资产在线管理后台 (SPA)。"""
    return HTMLResponse(get_admin_html())


# ==================== 管理后台 API 接口 ====================

# 内存元数据与页数轻量级缓存：file_path -> (mtime, pages_count, title)
_PPTX_META_CACHE: dict[str, tuple[float, int, str]] = {}


def _resolve_pptx_meta(file_path: str, mtime: float, fname: str, store_rec: dict | None = None) -> tuple[int, str]:
    """快速探测 PPTX 文件的总页数与标题，并在内存中基于 mtime 建立高命中缓存，避免重复打开文件。"""
    stem = os.path.splitext(fname)[0]

    # 1. 优先命中内存缓存（文件 mtime 未变动）
    cached = _PPTX_META_CACHE.get(file_path)
    if cached and cached[0] == mtime:
        return cached[1], cached[2]

    # 2. 标题提取：优先从 store spec 提取
    title = fname
    spec = store_rec.get("spec") if store_rec else None
    if spec and isinstance(spec, dict):
        title = spec.get("meta", {}).get("title") or spec.get("title") or fname

    # 3. 页数提取：优先检查 preview 截图数量（最快，0 次解析）
    pages_count = 1
    pv_dir = os.path.join(M.OUTPUT_DIR, "preview", M._slug(stem, "deck"))
    existing_pngs = glob.glob(os.path.join(pv_dir, "page_*.png"))
    if existing_pngs:
        pages_count = len(existing_pngs)
    elif spec and isinstance(spec, dict) and spec.get("pages"):
        pages_count = len(spec["pages"])
    else:
        # 仅当没有截图也没有 spec 时才轻度解析单文件
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            pages_count = len(prs.slides)
            if title == fname and prs.slides:
                for shape in prs.slides[0].shapes:
                    if shape.has_text_frame and shape.text_frame.text.strip():
                        first_line = shape.text_frame.text.strip().split("\n")[0]
                        if len(first_line) <= 50:
                            title = first_line
                        break
        except Exception:
            pages_count = 1

    # 写入缓存
    _PPTX_META_CACHE[file_path] = (mtime, pages_count, title)
    return pages_count, title


@app.post("/api/admin/verify")
def api_admin_verify(
    payload: VerifyIn | None = None,
    authorization: str | None = Header(None),
):
    """校验管理员访问凭证。"""
    if not AUTH_KEY:
        return {"ok": True, "auth_required": False}

    candidate = (payload.key if payload else "") or ""
    if not candidate and authorization:
        parts = authorization.strip().split()
        if len(parts) == 2 and parts[0].lower() == "bearer":
            candidate = parts[1]
        elif len(parts) == 1:
            candidate = parts[0]

    if candidate == AUTH_KEY:
        return {"ok": True, "auth_required": True}
    raise HTTPException(status_code=401, detail="Invalid AUTH_KEY")


@app.get("/api/admin/files")
def api_admin_files(
    authorization: str | None = Header(None),
    auth_key: str | None = Query(None),
    page: int = Query(1, ge=1, description="当前页码（从 1 开始）"),
    page_size: int = Query(15, ge=1, le=200, description="每页显示条数"),
    search: str = Query("", description="搜索关键词（匹配文件名或标题）"),
):
    """获取已生成的 PPTX 文稿资产列表（严格按生成/修改时间倒序排列，支持服务端毫秒级分页与极速缓存）。"""
    if AUTH_KEY and not check_auth(authorization=authorization, auth_key=auth_key):
        raise HTTPException(status_code=401, detail="Unauthorized: invalid or missing AUTH_KEY")

    store_items = {it["name"]: it for it in M.STORE.list(kind="pptx", limit=10000)}

    # 1. 毫秒级极速扫描磁盘：利用 os.scandir 只提取文件名、最后修改时间与大小，绝不打开文件
    scanned_files = []
    total_disk_bytes = 0

    if os.path.exists(M.OUTPUT_DIR):
        try:
            with os.scandir(M.OUTPUT_DIR) as it:
                for entry in it:
                    if entry.is_file() and entry.name.lower().endswith(".pptx"):
                        try:
                            st = entry.stat()
                            scanned_files.append({
                                "filename": entry.name,
                                "path": entry.path,
                                "mtime": st.st_mtime,
                                "size_bytes": st.st_size,
                            })
                            total_disk_bytes += st.st_size
                        except OSError:
                            continue
        except OSError:
            pass

    # 补充 STORE 中可能位于其他路径的条目
    scanned_names = {f["filename"] for f in scanned_files}
    for rec in store_items.values():
        fname = rec["name"]
        if fname not in scanned_names and os.path.exists(rec.get("path", "")):
            try:
                st = os.stat(rec["path"])
                scanned_files.append({
                    "filename": fname,
                    "path": rec["path"],
                    "mtime": st.st_mtime,
                    "size_bytes": st.st_size,
                })
                total_disk_bytes += st.st_size
                scanned_names.add(fname)
            except OSError:
                continue

    # 2. 严格按最后修改时间倒序排序（最新生成的在前）
    scanned_files.sort(key=lambda x: x["mtime"], reverse=True)

    # 3. 关键词过滤（若有搜索项，在内存中快速过滤）
    search_term = search.strip().lower()
    if search_term:
        filtered = []
        for f in scanned_files:
            fname = f["filename"]
            rec = store_items.get(fname)
            cached = _PPTX_META_CACHE.get(f["path"])
            cached_title = cached[2] if cached else ""
            store_title = (rec.get("spec", {}).get("meta", {}).get("title") or "") if rec else ""
            if search_term in fname.lower() or search_term in cached_title.lower() or search_term in store_title.lower():
                filtered.append(f)
        scanned_files = filtered

    total = len(scanned_files)
    import math
    total_pages = max(1, math.ceil(total / page_size)) if total > 0 else 1
    current_page = min(page, total_pages)

    # 4. 关键性能优化：仅对当前分页切片（如前 15~20 条）做详细页数与标题解析！
    start_idx = (current_page - 1) * page_size
    end_idx = start_idx + page_size
    current_slice = scanned_files[start_idx:end_idx]

    items = []
    for f in current_slice:
        fname = f["filename"]
        file_path = f["path"]
        mtime = f["mtime"]
        size_bytes = f["size_bytes"]
        stem = os.path.splitext(fname)[0]

        rec = store_items.get(fname)
        item_id = rec["id"] if rec else stem

        pages_count, title = _resolve_pptx_meta(file_path, mtime, fname, store_rec=rec)
        created_at = time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(mtime))

        items.append({
            "id": item_id,
            "filename": fname,
            "title": title,
            "pages_count": pages_count,
            "size_bytes": size_bytes,
            "mtime": mtime,
            "created_at": created_at,
            "thumbnail_url": f"/api/preview/{item_id}?page=1",
            "download_url": f"/api/download/{fname}",
        })

    return {
        "ok": True,
        "auth_required": bool(AUTH_KEY),
        "total": total,
        "total_disk_bytes": total_disk_bytes,
        "page": current_page,
        "page_size": page_size,
        "total_pages": total_pages,
        "items": items,
    }


@app.post("/api/admin/delete")
def api_admin_batch_delete(
    payload: BatchDeleteIn,
    authorization: str | None = Header(None),
    auth_key: str | None = Query(None),
):
    """批量彻底删除 PPTX 文稿资产及其预览截图缓存。"""
    if AUTH_KEY and not check_auth(authorization=authorization, auth_key=auth_key):
        raise HTTPException(status_code=401, detail="Unauthorized: invalid or missing AUTH_KEY")

    deleted_count = 0
    deleted_ids = []
    for item_id in payload.ids:
        if _delete_presentation_and_preview(item_id):
            deleted_count += 1
            deleted_ids.append(item_id)

    return {
        "ok": True,
        "deleted_count": deleted_count,
        "deleted_ids": deleted_ids,
    }


# ==================== 核心业务与 REST 端点 ====================

@app.get("/")
def index():
    return {
        "service": "pptx-studio",
        "version": "0.2.0",
        "auth_required": bool(AUTH_KEY),
        "admin_ui": "/admin",
        "mcp_endpoint": "/mcp",
        "docs": "/docs",
        "endpoints": [
            "GET /admin", "GET /manage", "GET /api/admin/files",
            "POST /api/admin/verify", "POST /api/admin/delete",
            "GET /healthz", "GET /api/themes", "GET /api/templates",
            "GET /api/templates/{name}", "POST /api/validate",
            "POST /api/generate", "POST /api/artifacts/{id}/page/{page_number}",
            "GET /api/artifacts", "GET /api/artifacts/{id}",
            "GET /api/download/{id}", "GET /api/preview/{id}",
            "DELETE /api/artifacts/{id}", "GET /api/server-info"
        ],
        "mcp_tools": [
            "get_design_guide", "list_themes", "list_templates", "get_template",
            "validate_spec", "create_presentation", "update_page", "change_theme",
            "insert_page", "delete_page", "render_preview", "list_artifacts",
            "read_spec", "delete_artifact", "server_info"
        ],
    }


@app.get("/healthz")
def healthz():
    return {
        "ok": True,
        "auth_required": bool(AUTH_KEY),
        "preview_available": M.render.available(),
        "artifacts": len(M.STORE.list(limit=10000)),
        "base_url": M.get_base_url(),
    }


@app.get("/api/themes")
def api_themes():
    return {
        "themes": [
            {
                "name": n,
                "label": p["label"],
                "mode": p.get("mode", "light"),
                "primary": p["primary"],
                "accent": p["accent"],
                "accent2": p.get("accent2", p["accent"]),
                "bg": p["bg"],
            }
            for n, p in theme_mod.PALETTES.items()
        ],
        "sizes": list(theme_mod.SLIDE_SIZES),
    }


@app.get("/api/templates")
def api_templates():
    return {"templates": templates.listing()}


@app.get("/api/templates/{name}")
def api_template(name: str):
    try:
        return templates.get(name)
    except KeyError as e:
        raise HTTPException(status_code=404, detail=str(e))


@app.post("/api/validate")
def api_validate(payload: SpecIn):
    errs = M._validate_spec(payload.spec)
    return {"ok": not errs, "errors": errs}


@app.post("/api/generate")
def api_generate(
    payload: SpecIn,
    authorization: str | None = Header(None),
    auth_key: str | None = Query(None),
):
    if AUTH_KEY and not check_auth(authorization=authorization, auth_key=auth_key):
        raise HTTPException(status_code=401, detail="Unauthorized: invalid or missing AUTH_KEY")

    errs = M._validate_spec(payload.spec)
    if errs:
        raise HTTPException(status_code=422, detail=errs)

    name = M.make_unique_filename(payload.filename, fallback="deck")
    out_path = os.path.join(M.OUTPUT_DIR, name)
    try:
        deck, path, warnings = build_from_spec(
            {**payload.spec, "output": {"path": out_path}})
    except Exception as e:
        raise HTTPException(status_code=500,
                            detail="%s: %s" % (type(e).__name__, e))
    rec = M.STORE.register(name, "pptx", path,
                           spec=payload.spec) if payload.store else {
        "id": None, "name": name, "path": path}

    base_url = M.get_base_url()
    download_url = f"{base_url}/api/download/{rec['name']}"
    alt_download_url = f"{base_url}/api/download/{rec['id']}.pptx" if rec.get("id") else download_url

    pages = payload.spec.get("pages") or []
    body = {
        "ok": True,
        "id": rec["id"],
        "file": rec["name"],
        "path": rec["path"],
        "pages": len(pages),
        "download_url": download_url,
        "alt_download_url": alt_download_url,
        "markdown_download_link": f"[点击下载演示文稿 (PPTX)]({download_url})",
        "warnings": warnings,
        "preview_urls": [],
        "slides": [],
        "user_display_markdown": "",
    }
    if payload.preview and M.render.available():
        try:
            pv = os.path.join(os.path.dirname(path), "preview",
                              M._slug(os.path.splitext(name)[0], "deck"))
            pdf, pngs = M.render.render_pptx(path, out_dir=pv)
            preview_urls = [
                f"{base_url}/api/preview/{rec['id'] or rec['name']}?page={i + 1}"
                for i in range(len(pngs))
            ]
            body["preview_urls"] = preview_urls
            slides_info = []
            for i in range(len(pngs)):
                p_idx = i + 1
                p_meta = pages[i] if i < len(pages) else {}
                slides_info.append({
                    "page": p_idx,
                    "title": p_meta.get("title", f"第 {p_idx} 页"),
                    "layout": p_meta.get("layout", "bullets"),
                    "preview_url": preview_urls[i]
                })
            body["slides"] = slides_info
            body["user_display_markdown"] = M.make_display_markdown(download_url, rec["name"], slides_info)
        except Exception as e:
            body["preview_error"] = str(e)
    elif payload.preview:
        body["preview_error"] = "未检测到 LibreOffice"
    return body


@app.post("/api/artifacts/{item_id:path}/page/{page_number}")
def api_update_page(
    item_id: str,
    page_number: int,
    payload: PageUpdateIn,
    authorization: str | None = Header(None),
    auth_key: str | None = Query(None),
):
    if AUTH_KEY and not check_auth(authorization=authorization, auth_key=auth_key):
        raise HTTPException(status_code=401, detail="Unauthorized: invalid or missing AUTH_KEY")

    res = M.update_page(item_id, page_number, payload.page_spec)
    data = M.json.loads(res[0])
    if not data.get("ok"):
        raise HTTPException(status_code=400, detail=data.get("error", "更新失败"))
    return data


@app.get("/api/artifacts")
def api_artifacts(
    kind: str = "",
    limit: int = Query(25, ge=1, le=200),
    authorization: str | None = Header(None),
    auth_key: str | None = Query(None),
):
    if AUTH_KEY and not check_auth(authorization=authorization, auth_key=auth_key):
        raise HTTPException(status_code=401, detail="Unauthorized: invalid or missing AUTH_KEY")
    return {"items": M.STORE.list(kind=kind or None, limit=limit)}


@app.get("/api/artifacts/{item_id}")
def api_artifact(
    item_id: str,
    authorization: str | None = Header(None),
    auth_key: str | None = Query(None),
):
    if AUTH_KEY and not check_auth(authorization=authorization, auth_key=auth_key):
        raise HTTPException(status_code=401, detail="Unauthorized: invalid or missing AUTH_KEY")
    rec = M.STORE.get(item_id)
    if not rec:
        raise HTTPException(status_code=404, detail="产物不存在")
    return rec


def _safe_resolve_path(item_id: str) -> str | None:
    """防路径穿越安全解析：确保目标文件必须真实存在且严格位于 OUTPUT_DIR 或 STORE.root 内部。"""
    raw_id = str(item_id or "").strip()
    if not raw_id or ".." in raw_id:
        return None

    # 去除首尾斜杠，仅取基名防止非法路径穿越
    clean_id = os.path.basename(raw_id)
    if not clean_id:
        return None

    target_path = None
    try:
        target_path = M.STORE.path_of(clean_id)
    except Exception:
        pass

    if not target_path or not os.path.exists(target_path):
        cand = os.path.join(M.OUTPUT_DIR, clean_id)
        if not os.path.exists(cand) and not cand.endswith(".pptx"):
            cand += ".pptx"
        if os.path.exists(cand):
            target_path = cand

    if target_path and os.path.exists(target_path):
        abs_target = os.path.abspath(target_path)
        abs_out = os.path.abspath(M.OUTPUT_DIR)
        abs_store = os.path.abspath(M.STORE.root)
        if abs_target.startswith(abs_out) or abs_target.startswith(abs_store):
            return abs_target

    return None


@app.get("/api/download/{item_id:path}")
def api_download(item_id: str):
    path = _safe_resolve_path(item_id)
    if not path or not os.path.exists(path):
        raise HTTPException(status_code=404, detail="文件不存在或无权访问")

    media_type = (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        if path.lower().endswith(".pptx")
        else "application/octet-stream"
    )
    return FileResponse(path, filename=os.path.basename(path), media_type=media_type)


@app.get("/api/preview/{item_id:path}")
def api_preview(item_id: str, page: int = Query(1, ge=1)):
    path = _safe_resolve_path(item_id)
    if not path or not os.path.exists(path):
        raise HTTPException(status_code=404, detail="文件不存在或无权访问")

    if path.lower().endswith(".png"):
        return FileResponse(path, media_type="image/png", headers={"Cache-Control": "public, max-age=86400"})

    pv_dir = os.path.join(os.path.dirname(path), "preview", M._slug(os.path.splitext(os.path.basename(path))[0], "deck"))
    target_png = os.path.join(pv_dir, "page_%02d.png" % page)

    # 1. 命中已生成的高清预览截图（毫秒级极速直出）
    if os.path.exists(target_png):
        return FileResponse(target_png, media_type="image/png", headers={"Cache-Control": "public, max-age=86400"})

    # 2. 如果预览目录已有生成的页面截图，判断请求页码是否超出总页数
    existing_pngs = sorted(glob.glob(os.path.join(pv_dir, "page_*.png")))
    if existing_pngs:
        if page > len(existing_pngs):
            raise HTTPException(status_code=404, detail=f"页码超出范围（当前演示文稿共 {len(existing_pngs)} 页）")
        if os.path.exists(existing_pngs[page - 1]):
            return FileResponse(existing_pngs[page - 1], media_type="image/png", headers={"Cache-Control": "public, max-age=86400"})

    # 3. 仅当从未生成过截图时，才调用无头转换补齐生成
    try:
        _, pngs = M.render.render_pptx(path, out_dir=pv_dir)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

    if page > len(pngs):
        raise HTTPException(status_code=404, detail="页码超出范围")

    return FileResponse(pngs[page - 1], media_type="image/png", headers={"Cache-Control": "public, max-age=86400"})


@app.delete("/api/artifacts/{item_id}")
def api_delete(
    item_id: str,
    authorization: str | None = Header(None),
    auth_key: str | None = Query(None),
):
    if AUTH_KEY and not check_auth(authorization=authorization, auth_key=auth_key):
        raise HTTPException(status_code=401, detail="Unauthorized: invalid or missing AUTH_KEY")

    if not _delete_presentation_and_preview(item_id):
        raise HTTPException(status_code=404, detail="产物不存在或已被删除")
    return {"ok": True, "deleted": item_id}


@app.get("/api/server-info")
def api_server_info():
    return JSONResponse(M.json.loads(M.server_info()))
